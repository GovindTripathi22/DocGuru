"""PDF Reference Mode DOCX base generator (VAL-04)."""

import logging
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

from ..models.template_spec import TemplateSpecification

logger = logging.getLogger(__name__)


class PdfReferenceBuilder:
    """
    Builds a fresh DOCX base that applies the geometry, margins,
    dominant fonts, and heading typography analyzed from a PDF template.
    """

    def build_docx_base(self, spec: TemplateSpecification, target_path: Path) -> Document:
        doc = Document()
        sec = doc.sections[0]

        # Apply page dimensions and margins
        if spec.page:
            sec.page_width = Pt(spec.page.width_pt)
            sec.page_height = Pt(spec.page.height_pt)
            sec.top_margin = Pt(spec.page.margins.top_pt)
            sec.bottom_margin = Pt(spec.page.margins.bottom_pt)
            sec.left_margin = Pt(spec.page.margins.left_pt)
            sec.right_margin = Pt(spec.page.margins.right_pt)

        # Apply typography
        normal_style = doc.styles["Normal"]
        body_font = spec.fonts.body or spec.fonts.default or "Calibri"
        normal_style.font.name = body_font
        normal_style.font.size = Pt(11)

        # Add or configure Heading styles
        for h_num in [1, 2, 3]:
            h_name = f"Heading {h_num}"
            try:
                h_style = doc.styles[h_name]
            except KeyError:
                h_style = doc.styles.add_style(h_name, WD_STYLE_TYPE.PARAGRAPH)

            h_font = getattr(spec.fonts, f"heading{h_num}", None) or body_font
            h_style.font.name = h_font
            h_style.font.bold = True
            sizes = {1: 18, 2: 14, 3: 12}
            h_style.font.size = Pt(sizes.get(h_num, 12))

        # Add headers / footers if text preview was detected across pages
        if spec.header and spec.header.has_content and spec.header.text_preview:
            sec.header.paragraphs[0].text = spec.header.text_preview

        if spec.footer and spec.footer.has_content and spec.footer.text_preview:
            sec.footer.paragraphs[0].text = spec.footer.text_preview

        target_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(target_path))
        return doc
