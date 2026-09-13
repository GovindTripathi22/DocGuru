from typing import Dict, List, Any, Optional, Set
import hashlib
import json
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from ..models.template_spec import (
    TemplateSpecification,
    SlideDimensionSpec,
    FontSpec,
    LayoutSpec,
    PlaceholderSpec,
    MasterSpec,
    ThemeSpec,
    LockStatus,
)

class PptxAnalyzer:
    """
    Deterministic template analyzer for PPTX files.
    Extracts slide masters, layouts, placeholders, dimensions, and typography.
    """

    def __init__(self):
        pass

    def analyze(self, file_path: str) -> TemplateSpecification:
        prs = Presentation(file_path)
        filename = Path(file_path).name

        # 1. Slide Dimensions
        slide_dims = self._extract_dimensions(prs)

        # 2. Slide Layouts & Placeholders
        layouts, layout_names = self._extract_layouts(prs)

        # 3. Slide Masters
        masters = self._extract_masters(prs)

        # 4. Fonts
        font_spec = self._extract_fonts(prs)

        # Calculate style hash
        style_hash = self._compute_pptx_style_hash(slide_dims, layouts, masters)

        return TemplateSpecification(
            filename=filename,
            document_type="pptx",
            total_pages_or_slides=len(prs.slides) if prs.slides else 1,
            slide_dimensions=slide_dims,
            fonts=font_spec,
            layouts=layouts,
            masters=masters,
            available_layout_names=layout_names,
            style_hash=style_hash,
            lock_status=LockStatus()
        )

    def _extract_dimensions(self, prs: Presentation) -> SlideDimensionSpec:
        width_in = prs.slide_width.inches if prs.slide_width else 13.333
        height_in = prs.slide_height.inches if prs.slide_height else 7.5
        ratio = "16:9"
        if abs((width_in / height_in) - (4.0 / 3.0)) < 0.1:
            ratio = "4:3"
        elif abs((width_in / height_in) - (16.0 / 9.0)) < 0.1:
            ratio = "16:9"

        return SlideDimensionSpec(
            width_inches=round(width_in, 3),
            height_inches=round(height_in, 3),
            aspect_ratio=ratio
        )

    def _extract_layouts(self, prs: Presentation):
        layouts_list: List[LayoutSpec] = []
        layout_names: List[str] = []

        for idx, layout in enumerate(prs.slide_layouts):
            placeholders: List[PlaceholderSpec] = []
            for p_idx, shape in enumerate(layout.placeholders):
                p_type = "UNKNOWN"
                try:
                    p_type = str(shape.placeholder_format.type).split(".")[-1]
                except Exception:
                    pass

                placeholders.append(PlaceholderSpec(
                    index=p_idx,
                    name=shape.name,
                    type=p_type,
                    left_inches=round(shape.left.inches if shape.left else 0.0, 2),
                    top_inches=round(shape.top.inches if shape.top else 0.0, 2),
                    width_inches=round(shape.width.inches if shape.width else 0.0, 2),
                    height_inches=round(shape.height.inches if shape.height else 0.0, 2)
                ))

            layout_spec = LayoutSpec(
                index=idx,
                name=layout.name,
                placeholders=placeholders
            )
            layouts_list.append(layout_spec)
            layout_names.append(layout.name)

        return layouts_list, layout_names

    def _extract_masters(self, prs: Presentation) -> List[MasterSpec]:
        masters: List[MasterSpec] = []
        for idx, master in enumerate(prs.slide_masters):
            masters.append(MasterSpec(
                index=idx,
                name=f"Master {idx + 1}",
                layout_count=len(master.slide_layouts)
            ))
        return masters

    def _extract_fonts(self, prs: Presentation) -> FontSpec:
        all_fonts: Set[str] = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        if p.font and p.font.name:
                            all_fonts.add(p.font.name)
                        for r in p.runs:
                            if r.font and r.font.name:
                                all_fonts.add(r.font.name)

        return FontSpec(
            default="Calibri",
            body="Calibri",
            heading1="Calibri Light",
            all_detected_fonts=sorted(list(all_fonts))
        )

    def _compute_pptx_style_hash(
        self,
        dims: SlideDimensionSpec,
        layouts: List[LayoutSpec],
        masters: List[MasterSpec]
    ) -> str:
        payload = {
            "dimensions": dims.model_dump(),
            "layout_count": len(layouts),
            "layout_names": [l.name for l in layouts],
            "master_count": len(masters)
        }
        raw_json = json.dumps(payload, sort_keys=True)
        return hashlib.sha256(raw_json.encode('utf-8')).hexdigest()[:16]
