import asyncio
import json
from pathlib import Path
from backend.app.analyzer.docx_analyzer import DocxAnalyzer
from backend.app.ai.planner import DocumentPlanner
from backend.app.ai.executor import DocumentExecutor
from backend.app.validation.diff_validator import DiffValidator
import docx
from docx import Document

async def run_live_pipeline():
    # 1. Create a branded sample template
    template_path = "backend/uploads/branded_company_template.docx"
    Path(template_path).parent.mkdir(parents=True, exist_ok=True)
    
    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "ACME CORP • CONFIDENTIAL RESEARCH DEPT"
    sec.footer.paragraphs[0].text = "Page 1 • Exact Branding Enforced"
    doc.add_paragraph("Template Title", style="Title")
    doc.add_paragraph("Original Heading 1", style="Heading 1")
    doc.add_paragraph("Original body text.")
    doc.save(template_path)

    # 2. Analyze template
    analyzer = DocxAnalyzer()
    spec = analyzer.analyze(template_path)
    print(f"[OK] Template Analyzed: {spec.filename} (Style Hash: {spec.style_hash})")

    # 3. Plan using local Gemma via Ollama
    planner = DocumentPlanner()
    print(">> Calling local Gemma model on your GPU to plan content...")
    plan = await planner.plan_document(
        user_prompt="Create a comprehensive technical report on AlphaGo breakthroughs and MCTS algorithm",
        template_spec=spec
    )
    print(f"[OK] Gemma Planned Document: '{plan.title}' with {len(plan.sections)} sections")

    # 4. Execute on template using pure application code (StyleLock enforced)
    executor = DocumentExecutor()
    output_path = "backend/outputs/generated_alphago_report.docx"
    executor.execute_docx(template_path, spec, plan, output_path)
    print(f"[OK] Document Generated at: {output_path}")

    # 5. Validate zero format drift
    validator = DiffValidator()
    val_report = validator.validate(spec, output_path)
    print(f"[OK] Diff Validation: Passed={val_report.passed}, Hash Match={val_report.hash_match}")
    print(f"[OK] Original Hash:  {val_report.original_style_hash}")
    print(f"[OK] Generated Hash: {val_report.generated_style_hash}")

    # 6. Verify content in generated file
    gen_doc = Document(output_path)
    print(f"[OK] Generated Document Header: '{gen_doc.sections[0].header.paragraphs[0].text}'")
    print(f"[OK] Generated Document Footer: '{gen_doc.sections[0].footer.paragraphs[0].text}'")
    first_heading = [p.text for p in gen_doc.paragraphs if p.style.name.startswith("Heading")][0]
    print(f"[OK] First Generated Heading: '{first_heading}'")

if __name__ == "__main__":
    asyncio.run(run_live_pipeline())
