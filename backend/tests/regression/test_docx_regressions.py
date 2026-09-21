"""Regression test suite for DOCX engine invariants (DOCX-01 to DOCX-05)."""

import os
import subprocess
import sys
from pathlib import Path
import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from backend.app.analyzer.template_analyzer import TemplateAnalyzer
from backend.app.engines.docx_engine import DocxEngine
from backend.app.errors import AppError
from backend.app.models.generation import DocumentPlan, DocumentSection, TableData, ThemeOverrideRequest


def test_create_document_contains_the_requested_title_and_conclusion(minimal_docx, tmp_path):
    """DOCX-01 & A-09: Create mode always renders title and conclusion."""
    spec = TemplateAnalyzer().analyze(str(minimal_docx))
    output = tmp_path / "output.docx"
    plan = DocumentPlan(title="T1", sections=[], conclusion="C1")

    engine = DocxEngine()
    engine.generate_from_plan(engine.load_template(str(minimal_docx), spec), plan, str(output))

    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "T1" in text
    assert "C1" in text
    assert "Conclusion" in text


def test_bullet_determinism_across_hash_seeds(minimal_docx, tmp_path):
    """DOCX-02 & A-10: Generation under multiple PYTHONHASHSEED values produces identical bullet styles."""
    script = f"""
import os
from docx import Document
from backend.app.analyzer.template_analyzer import TemplateAnalyzer
from backend.app.engines.docx_engine import DocxEngine
from backend.app.models.generation import DocumentPlan, DocumentSection

spec = TemplateAnalyzer().analyze(r"{minimal_docx}")
engine = DocxEngine()
plan = DocumentPlan(
    title="Determinism Test",
    sections=[DocumentSection(title="Sec 1", paragraphs=[], bullets=["Bullet A", "Bullet B"])]
)
out = r"{tmp_path}" + f"/bullet_{{os.environ.get('PYTHONHASHSEED', '0')}}.docx"
engine.generate_from_plan(engine.load_template(r"{minimal_docx}", spec), plan, out)
doc = Document(out)
styles = [p.style.name for p in doc.paragraphs if p.text in ("Bullet A", "Bullet B")]
print(",".join(styles))
"""
    results = set()
    for seed in ["0", "42", "1337", "9999"]:
        env = os.environ.copy()
        env["PYTHONHASHSEED"] = seed
        res = subprocess.run(
            [sys.executable, "-c", script],
            capture_output=True,
            text=True,
            env=env,
        )
        assert res.returncode == 0, f"Failed with seed {seed}: {res.stderr}"
        results.add(res.stdout.strip())

    assert len(results) == 1, f"Non-deterministic bullet styles observed across seeds: {results}"


def test_cover_table_preservation_and_run_formatting(tmp_path):
    """DOCX-03 & A-11: Cover table with bold/red/22pt title is preserved during placeholder fill."""
    doc_path = tmp_path / "cover_table.docx"
    doc = Document()
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    p = cell.paragraphs[0]
    run = p.add_run("{{title}}")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(255, 0, 0)
    doc.add_paragraph("Body paragraph to be inhabited.")
    doc.save(doc_path)

    spec = TemplateAnalyzer().analyze(str(doc_path))
    engine = DocxEngine()
    out = tmp_path / "filled_cover.docx"
    plan = DocumentPlan(title="Inhabited Report", sections=[DocumentSection(title="Sec 1", paragraphs=["New content"])])

    engine.generate_from_plan(engine.load_template(str(doc_path), spec), plan, str(out), template_mode="fill")

    res_doc = Document(out)
    assert len(res_doc.tables) == 1
    cell_p = res_doc.tables[0].cell(0, 0).paragraphs[0]
    assert "Inhabited Report" in cell_p.text
    assert cell_p.runs[0].bold is True


def test_edit_mode_target_not_found_raises_422(minimal_docx, tmp_path):
    """DOCX-04: Edit mode with nonexistent target heading raises 422 EDIT_TARGET_NOT_FOUND."""
    spec = TemplateAnalyzer().analyze(str(minimal_docx))
    engine = DocxEngine()
    plan = DocumentPlan(title="Edit", sections=[DocumentSection(title="New", paragraphs=["P"])])

    with pytest.raises(AppError) as exc_info:
        engine.generate_from_plan(
            engine.load_template(str(minimal_docx), spec),
            plan,
            str(tmp_path / "out.docx"),
            edit_action="insert_after",
            target_heading="Nonexistent Chapter 999",
        )
    assert exc_info.value.code == "EDIT_TARGET_NOT_FOUND"
    assert exc_info.value.http_status == 422


def test_theme_override_color_validation():
    """DOCX-05 & A-13: Invalid override color values fail validation; valid names map to hex."""
    with pytest.raises(ValueError):
        ThemeOverrideRequest(target="heading_color", value="invalid_not_a_color")

    ov = ThemeOverrideRequest(target="heading_color", value="teal")
    assert ov.value == "008080"
