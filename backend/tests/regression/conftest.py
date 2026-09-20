from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from pptx import Presentation


@pytest.fixture
def minimal_docx(tmp_path: Path) -> Path:
    """A programmatic template: regression tests must not add binary fixtures."""
    path = tmp_path / "template.docx"
    document = Document()
    document.sections[0].left_margin = Inches(1)
    document.add_paragraph("Template title", style="Title")
    document.add_paragraph("Template body", style="Normal")
    document.save(path)
    return path


@pytest.fixture
def minimal_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(path)
    return path
