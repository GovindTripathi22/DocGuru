"""Sensitivity tests asserting that DiffValidator detects tampering across every template component (VAL-01, VAL-02)."""

from pathlib import Path
import pytest
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation

from backend.app.analyzer.template_analyzer import TemplateAnalyzer
from backend.app.validation.diff_validator import DiffValidator
from backend.app.validation.fingerprint import fingerprint_package


def test_validator_idempotence_load_save_without_edits(minimal_docx, tmp_path):
    """VAL-02: Loading and saving a document without edits produces identical component fingerprints."""
    saved_doc = tmp_path / "idempotent.docx"
    doc = Document(minimal_docx)
    doc.save(saved_doc)

    orig_fp = fingerprint_package(minimal_docx)
    saved_fp = fingerprint_package(saved_doc)

    for comp, h in orig_fp.components.items():
        assert saved_fp.components.get(comp) == h, f"Component {comp} fingerprint drifted on load/save!"


def test_heading_style_size_and_color_mutation_fails_validation(minimal_docx, tmp_path):
    """VAL-01: Heading style font size/color mutations fail validation and name 'styles'."""
    original = TemplateAnalyzer().analyze(str(minimal_docx))
    mutated = tmp_path / "mutated_heading.docx"
    doc = Document(minimal_docx)
    style = doc.styles["Title"]
    style.font.size = Pt(48)
    style.font.color.rgb = RGBColor(255, 0, 0)
    doc.save(mutated)

    report = DiffValidator().validate(original, str(mutated), original_file_path=str(minimal_docx))
    assert report.passed is False
    assert report.status == "failed"
    assert any(c.name == "styles" and c.status == "changed" for c in report.components)


def test_margin_and_page_setup_mutation_fails_validation(minimal_docx, tmp_path):
    """VAL-01: Margin or page dimension mutation fails validation and names 'page_setup'."""
    original = TemplateAnalyzer().analyze(str(minimal_docx))
    mutated = tmp_path / "mutated_margins.docx"
    doc = Document(minimal_docx)
    sec = doc.sections[0]
    sec.top_margin = Inches(2.5)
    sec.left_margin = Inches(2.5)
    doc.save(mutated)

    report = DiffValidator().validate(original, str(mutated), original_file_path=str(minimal_docx))
    assert report.passed is False
    assert report.status == "failed"
    assert any(c.name == "page_setup" and c.status == "changed" for c in report.components)


def test_header_footer_mutation_fails_validation(sample_docx_template, tmp_path):
    """VAL-01: Header text alteration fails validation and names 'headers_footers'."""
    original = TemplateAnalyzer().analyze(str(sample_docx_template))
    mutated = tmp_path / "mutated_header.docx"
    doc = Document(sample_docx_template)
    sec = doc.sections[0]
    if sec.header.paragraphs:
        sec.header.paragraphs[0].text = "TAMPERED HEADER TEXT"
    else:
        sec.header.paragraphs.add_paragraph("TAMPERED HEADER TEXT")
    doc.save(mutated)

    report = DiffValidator().validate(original, str(mutated), original_file_path=str(sample_docx_template))
    assert report.passed is False
    assert report.status == "failed"
    assert any(c.name == "headers_footers" and c.status == "changed" for c in report.components)


def test_pptx_slide_size_mutation_fails_validation(minimal_pptx, tmp_path):
    """VAL-01: PPTX slide dimension alteration fails validation and names 'slide_size'."""
    original = TemplateAnalyzer().analyze(str(minimal_pptx))
    mutated = tmp_path / "mutated_size.pptx"
    prs = Presentation(minimal_pptx)
    prs.slide_width = Inches(8.5)
    prs.slide_height = Inches(11.0)
    prs.save(mutated)

    report = DiffValidator().validate(original, str(mutated), original_file_path=str(minimal_pptx))
    assert report.passed is False
    assert report.status == "failed"
    assert any(c.name == "slide_size" and c.status == "changed" for c in report.components)
