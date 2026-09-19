import pytest
from pathlib import Path
from docx import Document
from docx.shared import Inches

from backend.app.analyzer.docx_analyzer import DocxAnalyzer
from backend.app.engines.docx_engine import DocxEngine
from backend.app.ai.executor import DocumentExecutor
from backend.app.models.generation import DocumentPlan, DocumentSection, TableData
from backend.app.validation.diff_validator import DiffValidator

def test_extract_rules_and_outline(tmp_path: Path):
    """Test extracting formatting rules and outline from a structured document."""
    doc = Document()
    doc.add_paragraph("Seminar Report Format Guidelines", style="Title")
    doc.add_paragraph("Font Style: Times New Roman, Font Size: 12pt for body, 1.5 line spacing.")
    doc.add_paragraph("Margins: 1.25 inches for Left, 1.0 inch for Right, Top, and Bottom.")
    doc.add_paragraph("Citations must follow IEEE format: [1], [2].")
    doc.add_paragraph("1. INTRODUCTION", style="Heading 1")
    doc.add_paragraph("Introduction text goes here.")
    doc.add_paragraph("2. LITERATURE REVIEW", style="Heading 1")
    doc.add_paragraph("Review text goes here.")

    test_path = tmp_path / "rules_doc.docx"
    doc.save(str(test_path))

    analyzer = DocxAnalyzer()
    spec = analyzer.analyze(str(test_path))

    assert len(spec.extracted_rules) >= 3
    assert any("Times New Roman" in r or "12pt" in r for r in spec.extracted_rules)
    assert any("Margins" in r or "1.25" in r for r in spec.extracted_rules)
    assert any("IEEE" in r or "Citations" in r for r in spec.extracted_rules)
    assert "1. INTRODUCTION" in spec.document_outline
    assert "2. LITERATURE REVIEW" in spec.document_outline

def test_apply_content_edits_append_chapter(tmp_path: Path):
    """Test appending a new major chapter to an existing document without breaking original structure."""
    doc = Document()
    doc.add_paragraph("Cover Page: Annual AI Report", style="Title")
    doc.add_paragraph("1. Introduction", style="Heading 1")
    doc.add_paragraph("Existing introduction content.")
    
    template_path = tmp_path / "base_report.docx"
    doc.save(str(template_path))

    analyzer = DocxAnalyzer()
    spec = analyzer.analyze(str(template_path))

    engine = DocxEngine()
    locked = engine.load_template(str(template_path), spec)

    plan = DocumentPlan(
        title="Annual AI Report",
        edit_mode="edit",
        sections=[
            DocumentSection(
                title="5. Quantum Machine Learning Breakthroughs",
                heading_style="Heading 1",
                action="append",
                paragraphs=[
                    "Quantum Neural Networks leverage superposition and entanglement to accelerate tensor operations.",
                    "Recent experiments on 127-qubit systems achieve exponential speedups for specific eigenvalue problems."
                ],
                bullets=[
                    "Parameter shift rule enables gradient calculation on quantum circuits",
                    "Variational Quantum Eigensolver (VQE) demonstrates error mitigation tolerance"
                ],
                table_data=TableData(
                    headers=["Model Architecture", "Qubits Required", "Speedup vs Classical"],
                    rows=[
                        ["Q-CNN", "16", "3.2x"],
                        ["Quantum Transformer", "64", "14.8x"]
                    ],
                    title="Table 5.1: Comparative Performance Benchmark"
                )
            )
        ]
    )

    output_path = tmp_path / "appended_report.docx"
    engine.generate_from_plan(locked, plan, str(output_path))

    # Verify output document
    res_doc = Document(str(output_path))
    headings = [p.text for p in res_doc.paragraphs if p.text.startswith("5. Quantum")]
    assert len(headings) == 1
    assert "5. Quantum Machine Learning Breakthroughs" in headings[0]
    
    # Original text must still be present!
    assert any("Cover Page: Annual AI Report" in p.text for p in res_doc.paragraphs)
    assert any("Existing introduction content" in p.text for p in res_doc.paragraphs)
    
    # Verify table added
    assert len(res_doc.tables) >= 1
    table_text = " ".join(c.text for row in res_doc.tables[0].rows for c in row.cells)
    assert "Quantum Transformer" in table_text
    assert "14.8x" in table_text

def test_apply_content_edits_insert_after(tmp_path: Path):
    """Test inserting a new section directly after a specific heading landmark."""
    doc = Document()
    doc.add_paragraph("Report", style="Title")
    doc.add_paragraph("1. Introduction", style="Heading 1")
    doc.add_paragraph("Intro body.")
    doc.add_paragraph("2. Literature Review", style="Heading 1")
    doc.add_paragraph("Literature review body.")
    doc.add_paragraph("3. Conclusion", style="Heading 1")
    doc.add_paragraph("Conclusion body.")

    template_path = tmp_path / "base_insert.docx"
    doc.save(str(template_path))

    analyzer = DocxAnalyzer()
    spec = analyzer.analyze(str(template_path))

    engine = DocxEngine()
    locked = engine.load_template(str(template_path), spec)

    plan = DocumentPlan(
        title="Report",
        edit_mode="edit",
        sections=[
            DocumentSection(
                title="2.1 Modern Deep Learning Foundations",
                heading_style="Heading 1",
                action="insert_after",
                target_heading="2. Literature Review",
                paragraphs=["Inserted section analyzing transformer architectures and attention heads."]
            )
        ]
    )

    output_path = tmp_path / "inserted_report.docx"
    engine.generate_from_plan(locked, plan, str(output_path))

    res_doc = Document(str(output_path))
    paras = [p.text.strip() for p in res_doc.paragraphs if p.text.strip()]
    
    lit_idx = paras.index("2. Literature Review")
    insert_idx = paras.index("2.1 Modern Deep Learning Foundations")
    concl_idx = paras.index("3. Conclusion")

    # Inserted section must appear AFTER literature review and BEFORE conclusion
    assert lit_idx < insert_idx < concl_idx

def test_validation_passes_on_edited_docx(tmp_path: Path, sample_docx_template: Path):
    """Verify that editing/appending content preserves OpenXML style invariants and passes DiffValidator."""
    analyzer = DocxAnalyzer()
    spec = analyzer.analyze(str(sample_docx_template))

    plan = DocumentPlan(
        title="Corporate Research Template",
        edit_mode="edit",
        sections=[
            DocumentSection(
                title="2. Advanced Optimization Analysis",
                heading_style="Heading 1",
                action="append",
                paragraphs=["This new chapter adds comprehensive optimization analysis while honoring style lock."],
                bullets=["Convex approximation guarantees", "Stochastic gradient updates"]
            )
        ]
    )

    executor = DocumentExecutor()
    output_path = tmp_path / "validated_edit.docx"
    executor.execute_docx(str(sample_docx_template), spec, plan, str(output_path))

    validator = DiffValidator()
    report = validator.validate(spec, str(output_path))

    assert report.passed, f"Validation failed with issues: {report.issues}"
    assert report.hash_match
