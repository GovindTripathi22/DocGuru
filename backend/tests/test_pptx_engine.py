import pytest
from pathlib import Path
from pptx import Presentation
from backend.app.analyzer.pptx_analyzer import PptxAnalyzer
from backend.app.engines.pptx_engine import PptxEngine
from backend.app.models.generation import PresentationPlan, SlidePlan

def test_pptx_load_and_spec_extraction(sample_pptx_template: Path):
    analyzer = PptxAnalyzer()
    spec = analyzer.analyze(str(sample_pptx_template))

    assert spec.document_type == "pptx"
    assert spec.slide_dimensions is not None
    assert spec.slide_dimensions.aspect_ratio == "16:9"
    assert len(spec.available_layout_names) > 0
    assert len(spec.style_hash) > 0

def test_pptx_template_slide_generation(sample_pptx_template: Path, tmp_path: Path):
    analyzer = PptxAnalyzer()
    spec = analyzer.analyze(str(sample_pptx_template))
    engine = PptxEngine()

    locked_prs = engine.load_template(str(sample_pptx_template), spec)

    plan = PresentationPlan(
        presentation_title="AI Misinformation Presentation",
        audience="General Public",
        slides=[
            SlidePlan(
                slide_number=1,
                title="Combating AI Misinformation",
                subtitle="Strategies for Fact-Checking and Authenticity",
                layout_name="Title Slide",
                bullet_points=[]
            ),
            SlidePlan(
                slide_number=2,
                title="Key Challenges in Synthetic Media",
                layout_name="Title and Content",
                bullet_points=[
                    "Rapid generation of deepfakes and audio cloning",
                    "Difficulty in automated watermark detection",
                    "Erosion of public trust in digital communications"
                ]
            )
        ]
    )

    output_file = tmp_path / "generated_presentation.pptx"
    engine.generate_from_plan(locked_prs, plan, str(output_file))

    assert output_file.exists()

    # Verify generated presentation
    gen_prs = Presentation(str(output_file))
    assert len(gen_prs.slides) == 2

    # Check Slide 1 title and subtitle
    s1 = gen_prs.slides[0]
    assert s1.placeholders[0].text == "Combating AI Misinformation"
    assert s1.placeholders[1].text == "Strategies for Fact-Checking and Authenticity"

    # Check Slide 2 title and bullet content
    s2 = gen_prs.slides[1]
    assert s2.placeholders[0].text == "Key Challenges in Synthetic Media"
    body_text = s2.placeholders[1].text_frame.text
    assert "Rapid generation of deepfakes" in body_text
