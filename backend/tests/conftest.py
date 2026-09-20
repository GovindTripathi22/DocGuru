import pytest
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pptx
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt

from backend.app.config import settings

@pytest.fixture(scope="session", autouse=True)
def configure_test_provider():
    orig = settings.MODEL_PROVIDER
    if not settings.api_key:
        settings.MODEL_PROVIDER = "demo"
    yield
    settings.MODEL_PROVIDER = orig

@pytest.fixture(scope="session")
def fixtures_dir(tmp_path_factory) -> Path:
    base = tmp_path_factory.mktemp("fixtures")
    return base

@pytest.fixture(scope="session")
def sample_docx_template(fixtures_dir: Path) -> Path:
    """
    Creates a rich DOCX template with custom styles, headers, footers, borders, and a table.
    """
    file_path = fixtures_dir / "sample_report.docx"
    doc = Document()

    # Configure page setup
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Configure Header
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "ORGANIZATIONAL TEMPLATE: CONFIDENTIAL RESEARCH REPORT"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Configure Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "Page 1 of Document — All Rights Reserved"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add Title
    title = doc.add_paragraph("Corporate Research Template", style="Title")

    # Add Heading 1
    h1 = doc.add_paragraph("1. Section One Header", style="Heading 1")

    # Add Body Paragraph
    p1 = doc.add_paragraph("This is the original template body text formatted with standard line spacing and margins.")

    # Add Heading 2
    h2 = doc.add_paragraph("1.1 Sub-section Header", style="Heading 2")

    # Add a Styled Table
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    headers = ["Metric Name", "Baseline Score", "Target Benchmark"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        # Make header bold
        if cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].font.bold = True

    sample_data = [
        ["Accuracy", "92.4%", "98.5%"],
        ["Latency", "45ms", "12ms"]
    ]
    for r_idx, row in enumerate(sample_data):
        for c_idx, val in enumerate(row):
            table.cell(r_idx + 1, c_idx).text = val

    doc.save(str(file_path))
    return file_path

@pytest.fixture(scope="session")
def sample_pptx_template(fixtures_dir: Path) -> Path:
    """
    Creates a rich PPTX template with slide layouts and title/content placeholders.
    """
    file_path = fixtures_dir / "sample_presentation.pptx"
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    # Slide 1: Title Slide Layout (index 0)
    title_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_layout)
    slide1.placeholders[0].text = "Corporate Strategic Presentation Template"
    slide1.placeholders[1].text = "Q3 Executive Review & Operational Briefing"

    # Slide 2: Title and Content Layout (index 1)
    content_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(content_layout)
    slide2.placeholders[0].text = "Executive Summary & Objectives"
    tf = slide2.placeholders[1].text_frame
    tf.text = "Key priority 1: Maintain exact branding integrity"
    p2 = tf.add_paragraph()
    p2.text = "Key priority 2: Accelerate automated AI workflow adoption"

    prs.save(str(file_path))
    return file_path
